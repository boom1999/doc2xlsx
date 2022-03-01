# -*- coding: utf-8 -*-
# @Time : 2022/2/24 9:28
# @Author : lingz
# @Software: PyCharm

import copy
from docx import Document
import xlrd
import pandas as pd
import os


def get_docx_files_path(input_path):
    """
    input dirs of the docx_s, then listdir input_path-->dirs_path-->docx_files_path.
    :param input_path: the path containing the whole input files.
    :return: the path of the docx_file which need to tidy up.
    """
    docx_files_path = []
    docx_files = []
    dirs_path = os.listdir(input_path)
    for index in range(0, len(dirs_path)):
        dirs_path[index] = str(os.path.join(input_path, dirs_path[index]))
        files = os.listdir(dirs_path[index])
        for file in files:
            if os.path.splitext(file)[1] == ".docx":
                docx_files.append(file)
        docx_files_path.append(os.path.join(dirs_path[index]+'/', docx_files[index]))
    return docx_files_path


def read_docx(read_path):
    """

    :param read_path: .docx files.
    :return: list_ -- contain Dicts = {":", ";"...}.
    """
    input_path = read_path
    docx = Document(input_path)
    table_s = docx.tables

    # 从文件路径拆分出文件名，从文件名中按照‘-’拆出编码 code_
    (filepath, filename) = os.path.split(input_path)
    code_list = filename.split('-')[0:3]
    # TODO 换其他方法？
    code_ = code_list[0] + '-' + code_list[1] + '-' + code_list[2]

    list_ = []
    for table in table_s:
        for j in range(1, len(table.rows)):
            dict_ = {'序号': table.cell(j, 0).text,
                     '展品名称': table.cell(j, 1).text,
                     '所属领域': table.cell(j, 2).text,
                     '展品形式': table.cell(j, 3).text,
                     '展品单位': table.cell(j, 4).text,
                     '联系人': table.cell(j, 5).text,
                     '联系电话': table.cell(j, 6).text,
                     '编码': code_}
            if dict_['展品名称'] != '':
                list_.append(dict_)
                print("Append successfully: row " + str(j))
    return list_


def add_type_region(list_in, map_tables_path):
    """

    :param map_tables_path: path of the map, original list and their types and regions.
    :param list_in: putin original list which need to add type and region.
    :return: final list -- containing type and region.
    """

    # Use copy.deepcopy to avoid to change list_in outsider
    list_out = copy.deepcopy(list_in)
    # 这里需要拆分目录读取编码
    code_ = '编码'
    flat_ = '展品单位'
    flat_simple = '单位简称'
    type_ = '类型'
    origin_ = '地区'
    # Attention! 'map_tables_path' can't have head index, pd.read_excel can't read.
    # Or set 'header= '
    map_data = pd.read_excel(map_tables_path, index_col='编码', header=1, sheet_name='Sheet 1')
    code_list = list(map_data.index)
    for j in range(0, len(list_out)):
        if list_out[j][code_] not in code_list:
            print(list_out[j][flat_] + "出现编码错误，请手动修改!")
            continue
        list_out[j][flat_] = map_data.loc[list_out[j][code_], flat_]
        list_out[j][flat_simple] = map_data.loc[list_out[j][code_], flat_simple]
        list_out[j][type_] = map_data.loc[list_out[j][code_], type_]
        list_out[j][origin_] = map_data.loc[list_out[j][code_], origin_]

    return list_out


def export2excel(export, out):
    """

    :param export: list_sum -- contain Dicts = {":", ";"...} +  {":", ";"...} + ...
    :param out: outfile .xlsx file.
    :return:
    """
    pf = pd.DataFrame(list(export))

    # Redefine column labels
    order = ['编码', '序号', '展品名称', '所属领域', '展品形式', '展品单位', '单位简称', '类型', '地区', '联系人', '联系电话']
    pf = pf[order]

    file_path = pd.ExcelWriter(out)
    pf.fillna(' ', inplace=True)
    pf.to_excel(file_path, encoding='utf-8', index=False)
    file_path.save()


if __name__ == '__main__':
    # 输入word文件集目录
    input_path = '../docx/'
    # 输出excel表格目录
    output_path = "../xlsx/out_doc2xl.xlsx"
    # 匹配文件map目录
    map_tables_path = "../map/map_0301.xlsx"

    list_sum = []
    if os.path.exists(input_path):
        docx_files_path = get_docx_files_path(input_path)
        # test function: file_rename(input_path)
        # files = os.listdir(input_path)
        for docx_file_path in docx_files_path:
            tables = read_docx(input_path + str(docx_file_path))
            list_sum.extend(tables)
            print("Extend docx:' "+str(docx_file_path)+" 'successfully!\n")
    else:
        print('Path not exist')
    final_list = add_type_region(list_sum, map_tables_path)
    export2excel(final_list, output_path)
    print("Change successfully!")
